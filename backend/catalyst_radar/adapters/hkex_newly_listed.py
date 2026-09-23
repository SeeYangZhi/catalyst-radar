"""HKEX New Listing Report adapter (Phase 9.5b).

Confirms which HK IPOs have actually started trading. The HKEXnews
"New Listing Information" page links a per-year **New Listing Report**
workbook (``NLR<year>_Eng.xlsx``) carrying every Main Board listing
with its "Date of Listing" — a static file, no JS rendering needed
(the NLI HTML table itself has no listing dates). We download the
current year's workbook (plus the previous year's in January, so a
late-December listing isn't missed) and emit rows of
``{code, listing_date, name}`` with the stock code zero-padded to five
digits, matching the ``hkex:ipo:<code>`` Event key used by the
AAStocks discovery adapter.
"""

import io
from datetime import UTC, date, datetime
from typing import Any
from zoneinfo import ZoneInfo

import httpx

from catalyst_radar.adapters.base import FetchResult, SourceAdapter
from catalyst_radar.adapters.hk_codes import pad_code
from catalyst_radar.config import settings
from catalyst_radar.dedup import dedup_key as make_dedup_key
from catalyst_radar.dedup import source_event_id as make_source_event_id
from catalyst_radar.logging import get_logger

log = get_logger(__name__)


_HONG_KONG = ZoneInfo("Asia/Hong_Kong")


def _to_iso_date(value: Any) -> str | None:
    """Workbook 'Date of Listing' cell -> ISO date string."""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    s = str(value or "").strip()
    for fmt in ("%d/%m/%Y", "%d/%m/%y", "%Y-%m-%d"):
        try:
            return datetime.strptime(s, fmt).date().isoformat()
        except ValueError:
            continue
    return None


def parse_new_listing_report(xlsx_bytes: bytes) -> list[dict[str, Any]]:
    """Pure: NLR workbook bytes -> [{code, listing_date, name}].

    Scans each sheet for a header row carrying "Stock Code" and
    "Date of Listing", then emits one row per listing. Continuation
    rows (the funds-raised "(b)" lines repeat ``"`` in every cell) and
    malformed rows are skipped fail-soft; an unreadable workbook or a
    sheet without the expected header yields ``[]``.
    """
    from openpyxl import load_workbook

    try:
        wb = load_workbook(io.BytesIO(xlsx_bytes), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001 - corrupt/partial download
        log.warning("hkex_nlr_unreadable", error=repr(exc))
        return []

    out: list[dict[str, Any]] = []
    seen: set[str] = set()
    try:
        for ws in wb.worksheets:
            code_col: int | None = None
            date_col: int | None = None
            name_col: int | None = None
            for row in ws.iter_rows(values_only=True):
                if code_col is None or date_col is None:
                    for i, cell in enumerate(row):
                        low = str(cell or "").strip().lower()
                        if low.startswith("stock code"):
                            code_col = i
                        elif low.startswith("date of listing"):
                            date_col = i
                        elif low.startswith("company name"):
                            name_col = i
                    continue
                try:
                    code = pad_code(row[code_col])
                    if not code or code in seen:
                        continue
                    listing_date = _to_iso_date(row[date_col])
                    if not listing_date:
                        continue
                    name = ""
                    if name_col is not None and name_col < len(row):
                        name = " ".join(str(row[name_col] or "").split())
                    out.append({"code": code, "listing_date": listing_date, "name": name})
                    seen.add(code)
                except Exception:  # noqa: BLE001 - one bad row never aborts
                    continue
    finally:
        wb.close()
    return out


class HkexNewlyListedAdapter(SourceAdapter):
    source_name = "hkexnews"
    schema_name = "hkexnews.newly_listed.v1"

    def __init__(self, timeout: int | None = None) -> None:
        self.timeout = timeout or settings.hkex_request_timeout_seconds
        self.url_template = settings.hkex_new_listing_report_url_template

    def report_urls(self, today: date | None = None) -> list[str]:
        # The report is keyed to the Hong Kong calendar year, so "today"
        # must be the HK wall clock, not UTC.
        d = today or datetime.now(_HONG_KONG).date()
        years = [d.year]
        if d.month == 1:  # year-boundary: late-Dec listings live in last year's report
            years.append(d.year - 1)
        return [self.url_template.format(year=y) for y in years]

    async def fetch(self, target: Any = None) -> FetchResult:
        urls = self.report_urls()
        items: list[dict[str, Any]] = []
        statuses: list[int] = []
        seen: set[str] = set()
        try:
            async with httpx.AsyncClient(timeout=self.timeout, follow_redirects=True) as client:
                for url in urls:
                    try:
                        resp = await client.get(
                            url, headers={"User-Agent": settings.scraper_user_agent}
                        )
                    except httpx.HTTPError as exc:
                        log.warning("hkex_nlr_fetch_failed", url=url, error=repr(exc))
                        statuses.append(0)
                        continue
                    statuses.append(resp.status_code)
                    if resp.status_code != 200:
                        log.warning("hkex_nlr_http_error", url=url, status=resp.status_code)
                        continue
                    for row in parse_new_listing_report(resp.content):
                        if row["code"] in seen:
                            continue
                        seen.add(row["code"])
                        items.append(row)
        except Exception as exc:  # noqa: BLE001 - fetch never raises
            log.warning("hkex_nlr_failed", error=repr(exc))
        http_status = 200 if 200 in statuses else (statuses[0] if statuses else 0)
        return FetchResult(
            source_name=self.source_name,
            schema_name=self.schema_name,
            source_url=urls[0],
            http_status=http_status,
            payload={"urls": urls, "rows": items},
            items=items,
        )

    def normalize(self, raw_item: dict[str, Any]) -> dict[str, Any]:
        code = pad_code(raw_item.get("code")) or ""
        listing_date = raw_item.get("listing_date")
        return {
            "event_type": "ipo",
            "symbol": code,
            "exchange": "HKSE",
            "company_name": (raw_item.get("name") or "").strip() or None,
            "event_date": (
                datetime.fromisoformat(listing_date).replace(tzinfo=UTC)
                if listing_date
                else None
            ),
            "payload": {"code": code, "listing_date": listing_date},
        }

    def source_event_id(self, raw_item: dict[str, Any]) -> str:
        # Same key space as the AAStocks adapter so a report row resolves
        # to the existing hkex:ipo:<code> Event.
        return make_source_event_id("hkex", "ipo", pad_code(raw_item.get("code")))

    def dedup_key(self, raw_or_normalized_item: dict[str, Any]) -> str:
        return make_dedup_key(self.source_event_id(raw_or_normalized_item))
